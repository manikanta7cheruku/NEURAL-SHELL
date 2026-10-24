"""
The document library reads PDFs page by page, finds passages by keyword and meaning, keeps documents attached to
a chat, and gives the model formatting rules and page citations for study answers.
"""
import os
import tempfile
import unittest

from support import BrainTest, LLM_CALLS
import brain
from knowledge import core, store

_PAGES = [
    ["1 Introduction to Operating Systems", "An operating system manages hardware and software resources.",
     "It provides services such as process scheduling and memory management."],
    ["2 Process Scheduling", "First Come First Served runs processes in arrival order.",
     "Round Robin gives each process a fixed time quantum of 10 to 100 milliseconds.",
     "Shortest Job First picks the process with the smallest burst time."],
    ["3 Memory Management", "Paging divides memory into fixed size frames of 4 KB.",
     "Segmentation divides memory into variable size segments.",
     "A page fault occurs when a page is not in main memory."],
]


def make_pdf(path, pages):
    """A minimal text PDF, one page per list of lines."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        stream = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({line}) '" for line in lines) + " ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents {len(objs)} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    with open(path, "w", encoding="latin-1") as f:
        f.write(out)
    return path


class Library(BrainTest):
    def setUp(self):
        super().setUp()
        store.clear()
        self.dir = tempfile.mkdtemp()
        self.pdf = make_pdf(os.path.join(self.dir, "os_notes.pdf"), _PAGES)

    def add(self, path, conv=None):
        doc = store.wait(store.add(path, conversation_id=conv)["id"])
        self.assertEqual(doc["status"], "ready", doc.get("error"))
        return doc

    def last_prompt(self):
        return LLM_CALLS[-1]["messages"][-1]["content"]

    def test_pdf_is_read_with_pages_and_sections(self):
        doc = self.add(self.pdf)
        self.assertEqual(doc["unit"], "3 pages")
        self.assertEqual([h["title"] for h in doc["headings"]][:3],
                         ["1 Introduction to Operating Systems", "2 Process Scheduling", "3 Memory Management"])
        self.assertEqual(store.search("round robin time quantum")[0]["label"], "p. 2")

    def test_same_file_twice_is_one_document(self):
        self.assertEqual(self.add(self.pdf)["id"], self.add(self.pdf)["id"])
        self.assertEqual(len(store.documents()), 1)

    def test_attached_document_answers_every_question_in_that_chat(self):
        doc = self.add(self.pdf)
        brain.think("What is round robin?", "default", "text", self.conv, documents=[doc["id"]])
        prompt = self.last_prompt()
        self.assertIn("<knowledge_context>", prompt)
        self.assertIn("[p. 2", prompt)
        self.assertIn("markdown table", prompt)
        self.assertIn({"name": "os_notes.pdf", "label": "p. 2"}, brain.get_trace("default", self.conv)["citations"])
        self.ask("And what is a page fault")                     # no attachment sent: the chat remembers it
        self.assertIn("A page fault occurs", self.last_prompt())

    def test_summary_spreads_across_the_whole_document(self):
        doc = self.add(self.pdf, conv=self.conv)
        self.ask("Summarize it")
        prompt = self.last_prompt()
        for page in ("[p. 1", "[p. 2", "[p. 3"):
            self.assertIn(page, prompt)
        self.assertIn("Its sections:", prompt)
        self.assertTrue(doc["id"] in store.attached(self.conv))

    def test_unattached_library_only_used_when_it_matches(self):
        self.add(self.pdf)
        self.ask("Tell me a joke about cats")
        self.assertNotIn("<knowledge_context>", self.last_prompt())
        self.ask("Explain the difference between paging and segmentation")
        self.assertIn("Segmentation divides memory", self.last_prompt())

    def test_open_the_pdf_opens_the_attached_document(self):
        self.add(self.pdf, conv=self.conv)
        self.assertEqual(self.ask("open the pdf"), "Opened os_notes.pdf.")

    def test_unreadable_files_say_why(self):
        empty = make_pdf(os.path.join(self.dir, "scan.pdf"), [[], []])
        doc = store.wait(store.add(empty)["id"])
        self.assertEqual(doc["status"], "failed")
        self.assertIn("scanned", doc["error"])

    def test_remove_deletes_the_copy_and_passages(self):
        doc = self.add(self.pdf, conv=self.conv)
        path = store.get(doc["id"])["path"]
        self.assertTrue(store.remove(doc["id"]))
        self.assertFalse(os.path.exists(path))
        self.assertEqual(store.attached(self.conv), [])
        self.assertEqual(store.search("round robin"), [])


class Formatting(unittest.TestCase):
    def test_budget_grows_with_the_context_window(self):
        self.assertLess(core.budget_chars(4096), core.budget_chars(8192))

    def test_whole_document_questions(self):
        for q in ("summarize this", "make me revision notes", "what is this document about", "key points"):
            self.assertTrue(store.wants_whole_document(q), q)
        self.assertFalse(store.wants_whole_document("what is a page fault"))

    def test_word_tables_and_headings(self):
        import importlib.util
        if importlib.util.find_spec("docx") is None:
            self.skipTest("python-docx not installed")
        import docx
        from knowledge import extract
        d = docx.Document()
        d.add_heading("Units", 1)
        table = d.add_table(rows=2, cols=2)
        for r, row in enumerate([["Quantity", "Unit"], ["Capacitance", "Farad"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        path = os.path.join(tempfile.mkdtemp(), "units.docx")
        d.save(path)
        text = extract.extract(path)[0]["text"]
        self.assertIn("# Units", text)
        self.assertIn("Capacitance | Farad", text)


if __name__ == "__main__":
    unittest.main()
