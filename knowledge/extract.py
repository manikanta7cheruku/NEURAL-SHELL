"""
knowledge/extract.py

Turns a document into units Seven can cite: one per PDF page, slide or sheet (a Word or text file is one unit,
its headings marked with "#"). Tables are kept as rows of cells separated by " | " so answers can rebuild them.
"""

import logging
import os
import re
from typing import Dict, List

SUPPORTED = (".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md")
_MAX_SHEET_ROWS = 5000
logging.getLogger("pypdf").setLevel(logging.ERROR)      # damaged-but-readable PDFs are common

# "Chapter 3 ...", "Unit 2: ...", "2.1 Process States", "4 Memory" are headings; "1. Buy milk" is a list item.
_HEADING_RE = re.compile(
    r"^(?:(?i:chapter|unit|module|section|part|lesson)\s+[\dIVX]+\b.*|\d{1,2}(?:\.\d{1,2}){1,2}\.?\s+[A-Z].{2,70}|"
    r"\d{1,2}\s+[A-Z][A-Za-z].{2,70})$")


class ExtractError(Exception):
    """A document Seven cannot read, with a reason written for the user."""


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not 3 <= len(line) <= 80 or line.endswith((".", ",", ";", ":")):
        return False
    if _HEADING_RE.match(line):
        return True
    words = line.split()
    return 2 <= len(words) <= 8 and line.isupper() and sum(len(w) >= 3 and w.isalpha() for w in words) >= 2


def _pdf(path: str) -> List[Dict]:
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractError("This PDF is password protected. Remove the password and add it again.")
        pages = reader.pages
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError(f"This PDF could not be opened ({exc.__class__.__name__}). It may be damaged.")
    units = []
    for n, page in enumerate(pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = re.sub(r"[ \t]+", " ", text).strip()
        if text:
            units.append({"text": text, "page": n, "label": f"p. {n}"})
    if len(pages) and sum(len(u["text"]) for u in units) < 40 * max(1, len(pages)) // 4:
        raise ExtractError("This PDF looks like scanned images, so there is no text to read. "
                           "Export it with text (OCR) and add it again.")
    return units


def _docx(path: str) -> List[Dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = docx.Document(path)
    lines = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "").lower()
            if style.startswith(("heading", "title")):
                lines.append(f"\n# {text}\n")
            else:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:      # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
    return [{"text": "\n".join(lines).strip(), "page": None, "label": None}]


def _pptx(path: str) -> List[Dict]:
    from pptx import Presentation
    units = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(" ".join(c.text.split()) for c in row.cells))
            elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            units.append({"text": "\n".join(parts), "page": n, "label": f"slide {n}"})
    return units


def _xlsx(path: str) -> List[Dict]:
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    units = []
    try:
        for name in book.sheetnames:
            rows = []
            for row in book[name].iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(" | ".join(cells))
                if len(rows) >= _MAX_SHEET_ROWS:
                    break
            if rows:
                units.append({"text": "\n".join(rows), "page": None, "label": f"sheet {name}"})
    finally:
        book.close()
    return units


def _text(path: str) -> List[Dict]:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return [{"text": f.read().strip(), "page": None, "label": None}]


def extract(path: str) -> List[Dict]:
    """Units of {"text", "page", "label"} in reading order. Raises ExtractError with a user-facing reason."""
    ext = os.path.splitext(path)[1].lower()
    readers = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".txt": _text, ".md": _text}
    if ext not in readers:
        raise ExtractError(f"{ext or 'This file type'} isn't supported. Use PDF, Word, PowerPoint, Excel, TXT or MD.")
    try:
        units = [u for u in readers[ext](path) if u["text"]]
    except ExtractError:
        raise
    except ImportError as exc:
        raise ExtractError(f"The reader for {ext} files is missing ({exc.name}). Reinstall Seven to restore it.")
    except Exception as exc:
        raise ExtractError(f"This file could not be read ({exc.__class__.__name__}). It may be damaged.")
    if not units:
        raise ExtractError("There is no readable text in this file.")
    return units


def headings(units: List[Dict], limit: int = 60) -> List[Dict]:
    """Likely section titles with where they appear, for outlines and the document view."""
    found, seen = [], set()
    for unit in units:
        for line in unit["text"].splitlines():
            line = line.strip()
            if line.startswith("#"):
                line = line.lstrip("#").strip()
            elif not _looks_like_heading(line):
                continue
            key = line.lower()
            if line and key not in seen:
                seen.add(key)
                found.append({"title": line, "label": unit["label"]})
                if len(found) >= limit:
                    return found
    return found
